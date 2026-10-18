"""Carga en Jano (Recetas) los ingredientes y sus vínculos a productos de Compras a partir del Excel
«Recetas - ingredientes por producto (propuesta).xlsx» revisado por Luis.

Uso:  python3 scripts/recetas-ingredientes/cargar.py "~/Downloads/Recetas - ingredientes por producto (propuesta).xlsx"

- Lee la hoja Productos: columna «Ingrediente (propuesta)» y, si está rellena, «OK / corrección» manda
  («NO» = sin ingrediente). Unidad de la columna «Unidad» (kg/l/ud; si falta, kg).
- Familia/subfamilia del ingrediente = la más frecuente entre sus productos (texto de Ratios).
- BORRA antes todos los vínculos e ingredientes de la cuenta que no estén en uso por líneas de receta
  (los que estén en uso se conservan y se reutilizan por nombre).
- Inserta ingredientes y vínculos (factor 1). El principal se marca después por SQL (más líneas de
  albarán en 12 meses) — ver el bloque final.
- Claves en apps/general/.env.local (service role). Idempotente: se puede repetir.
"""
import json, os, re, sys, unicodedata, urllib.request, uuid
from collections import Counter, defaultdict
from openpyxl import load_workbook

RAIZ = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
CUENTA = '082c5366-d9ae-49b9-a8b8-8caad73985bd'
env = {}
for l in open(os.path.join(RAIZ, 'apps/general/.env.local')):
    l = l.strip()
    if '=' in l and not l.startswith('#'):
        k, v = l.split('=', 1); env[k] = v.strip().strip('"').strip("'")
URL, KEY = env['NEXT_PUBLIC_SUPABASE_URL'], env['SUPABASE_SERVICE_ROLE_KEY']
H = {'apikey': KEY, 'Authorization': 'Bearer ' + KEY, 'Content-Type': 'application/json'}

def api(method, path, body=None, prefer=None, rango=None):
    h = dict(H)
    if prefer: h['Prefer'] = prefer
    if rango: h['Range-Unit'] = 'items'; h['Range'] = rango
    req = urllib.request.Request(URL + '/rest/v1/' + path, method=method, headers=h, data=json.dumps(body).encode() if body is not None else None)
    try:
        with urllib.request.urlopen(req) as r:
            t = r.read().decode(); return json.loads(t) if t else None
    except urllib.error.HTTPError as e:
        raise SystemExit(f'{method} {path} → {e.code}: {e.read().decode()[:400]}')

def fetch_all(path):
    out, desde = [], 0
    while True:
        chunk = api('GET', path, rango=f'{desde}-{desde+999}') or []
        out += chunk
        if len(chunk) < 1000: return out
        desde += 1000

def norm(s):
    s = unicodedata.normalize('NFD', (s or '').lower()); s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9%]+', ' ', s).strip()

# ── 1. Excel ───────────────────────────────────────────────────────────────
xlsx = os.path.expanduser(sys.argv[1] if len(sys.argv) > 1 else '~/Downloads/Recetas - ingredientes por producto (propuesta).xlsx')
ws = load_workbook(xlsx, read_only=True)['Productos']
filas = list(ws.iter_rows(values_only=True)); cab = [str(c or '').strip() for c in filas[0]]
col = {n: cab.index(n) for n in ['Id producto', 'Producto', 'Familia', 'Subfamilia', 'Ingrediente (propuesta)', 'Unidad', 'OK / corrección']}
asig = {}  # id_producto → (ingrediente, unidad, familia, subfamilia)
correcciones = 0
for f in filas[1:]:
    if not f or not f[col['Id producto']]: continue
    idp = str(f[col['Id producto']]).strip()
    ing = str(f[col['Ingrediente (propuesta)']] or '').strip()
    ok = str(f[col['OK / corrección']] or '').strip()
    if ok:
        correcciones += 1
        ing = '' if ok.upper() in ('NO', 'N', '-', '—') else ok
    ud = str(f[col['Unidad']] or '').strip().lower()
    if ud not in ('kg', 'l', 'ud'): ud = 'kg'
    if ing: asig[idp] = (re.sub(r'\s+', ' ', ing), ud, str(f[col['Familia']] or '').upper().strip(), str(f[col['Subfamilia']] or '').strip())
print(f'Excel: {len(filas)-1} filas, {len(asig)} con ingrediente, {correcciones} correcciones de Luis')

# ── 2. Productos de Compras (uuid por código) ──────────────────────────────
prods = {p['codigo_interno']: p['id'] for p in fetch_all(f'compras_producto?cuenta_id=eq.{CUENTA}&select=id,codigo_interno')}
sin_prod = [k for k in asig if k not in prods]
if sin_prod: print('AVISO: códigos del Excel que ya no están en Compras:', sin_prod[:10], '…' if len(sin_prod) > 10 else '')

# ── 3. Ingredientes a crear (nombre canónico: la grafía más frecuente) ─────
formas = defaultdict(Counter); datos = defaultdict(lambda: {'ud': Counter(), 'fam': Counter(), 'sub': Counter(), 'prods': []})
for idp, (ing, ud, fam, sub) in asig.items():
    if idp not in prods: continue
    formas[norm(ing)][ing] += 1
for idp, (ing, ud, fam, sub) in asig.items():
    if idp not in prods: continue
    n = norm(ing); d = datos[n]; d['ud'][ud] += 1; d['fam'][fam or ''] += 1; d['sub'][sub or ''] += 1; d['prods'].append(prods[idp])
nombre = {n: formas[n].most_common(1)[0][0] for n in formas}
print(f'Ingredientes distintos: {len(datos)}')

# ── 4. Borrar precarga anterior (conservando los ingredientes en uso) ──────
lineas = fetch_all(f'rec_receta_linea?cuenta_id=eq.{CUENTA}&select=ingrediente_id&ingrediente_id=not.is.null')
en_uso = {l['ingrediente_id'] for l in lineas}
existentes = fetch_all(f'rec_ingrediente?cuenta_id=eq.{CUENTA}&select=id,nombre,unidad')
api('DELETE', f'rec_ingrediente_producto?cuenta_id=eq.{CUENTA}', prefer='return=minimal')
borrables = [i['id'] for i in existentes if i['id'] not in en_uso]
for i in range(0, len(borrables), 100):
    api('DELETE', f'rec_ingrediente?cuenta_id=eq.{CUENTA}&id=in.({",".join(borrables[i:i+100])})', prefer='return=minimal')
conservados = {norm(i['nombre']): i for i in existentes if i['id'] in en_uso}
print(f'Borrados {len(borrables)} ingredientes de la precarga y todos sus vínculos; conservados en uso: {len(conservados)}')

# ── 5. Insertar ingredientes ───────────────────────────────────────────────
nuevos, ids = [], {}
for n, d in datos.items():
    if n in conservados:
        ids[n] = conservados[n]['id']; continue
    i = str(uuid.uuid4()); ids[n] = i
    nuevos.append({'id': i, 'cuenta_id': CUENTA, 'nombre': nombre[n], 'unidad': d['ud'].most_common(1)[0][0],
                   'familia': d['fam'].most_common(1)[0][0], 'subfamilia': d['sub'].most_common(1)[0][0], 'coste_modo': 'auto', 'activo': True})
for i in range(0, len(nuevos), 200):
    api('POST', 'rec_ingrediente', nuevos[i:i+200], prefer='return=minimal')
print(f'Insertados {len(nuevos)} ingredientes')

# ── 6. Insertar vínculos ───────────────────────────────────────────────────
vinc = [{'cuenta_id': CUENTA, 'ingrediente_id': ids[n], 'producto_id': p, 'factor': 1, 'principal': False} for n, d in datos.items() for p in d['prods']]
for i in range(0, len(vinc), 300):
    api('POST', 'rec_ingrediente_producto', vinc[i:i+300], prefer='return=minimal')
print(f'Insertados {len(vinc)} vínculos producto → ingrediente (factor 1)')
print('Ahora marcar el principal por SQL (ver bloque al final del fichero).')

# ── 7. Principal (ejecutar en SQL, no desde aquí) ──────────────────────────
PRINCIPAL_SQL = """
with n as (
  select v.id, v.ingrediente_id,
         (select count(*) from compras_linea l join compras_doc d on d.id = l.doc_id
           where l.producto_id = v.producto_id and d.fecha >= current_date - 365) as n_lineas,
         c.nombre
  from rec_ingrediente_producto v join compras_producto c on c.id = v.producto_id
  where v.cuenta_id = '082c5366-d9ae-49b9-a8b8-8caad73985bd'
), r as (
  select id, row_number() over (partition by ingrediente_id order by n_lineas desc, nombre) pos from n
)
update rec_ingrediente_producto v set principal = (r.pos = 1) from r where r.id = v.id;
"""
